from copy import copy
from datetime import date
import random

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SOURCE = "/Users/rd/Downloads/Dealer360_Specification_Model_Template.xlsx"
OUTPUT = "/Users/rd/Documents/Dealer360/Dealer360_Specification_Model_Sample_45_Dealers_3_Years.xlsx"

rng = random.Random(3602026)
wb = load_workbook(SOURCE)
ws = wb["Dealer_Data"]

states = [
    ("Maharashtra", "West"), ("Gujarat", "West"), ("Rajasthan", "North"),
    ("Delhi", "North"), ("Uttar Pradesh", "North"), ("Punjab", "North"),
    ("Karnataka", "South"), ("Tamil Nadu", "South"), ("Kerala", "South"),
    ("Telangana", "South"), ("Andhra Pradesh", "South"), ("West Bengal", "East"),
    ("Odisha", "East"), ("Bihar", "East"), ("Madhya Pradesh", "Central"),
    ("Chhattisgarh", "Central"), ("Assam", "North East"), ("Haryana", "North"),
]
prefixes = ["Apex", "Metro", "Prime", "Shree", "Sterling", "Crown", "Pioneer", "Summit", "Capital", "Evergreen", "Royal", "United", "Classic", "Velocity", "Horizon"]
suffixes = ["Motors", "Auto", "Wheels", "Mobility", "Automobiles", "Vehicles", "Drive"]

dealers = []
used = set()
for i in range(45):
    state, region = states[i % len(states)]
    while True:
        name = f"{prefixes[i % len(prefixes)]} {suffixes[(i * 3 + i // 7) % len(suffixes)]} {i+1:02d}"
        if name not in used:
            used.add(name)
            break
    dealers.append({
        "id": f"DLR-{i+1:03d}", "name": name, "state": state, "region": region,
        "target": rng.uniform(82, 108), "growth": rng.uniform(-2, 13),
        "inventory": rng.uniform(54, 116), "delay": rng.uniform(1, 10),
        "sla": rng.uniform(82, 98), "complaints": rng.uniform(0.8, 4.2),
        "potential": rng.randint(70, 460), "frt": rng.uniform(1.2, 6.0),
        "art": rng.uniform(8, 36), "satisfaction": rng.uniform(3.5, 4.7),
    })

headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
assert headers == [
    "Dealer ID", "Dealer Name", "Month", "State", "Region", "Sales versus Target (%)",
    "Sales Growth YoY (%)", "Inventory Ageing (days)", "Payment Delays (Days)",
    "Service Performance (SLA, %)", "Customer Complaints (per 100 Retail Sales)",
    "Local Market Potential (units/month)", "FRT (hours)", "ART (hours)",
    "Customer Satisfaction Score (0-5)",
]

# Apply an unobtrusive header style and preserve the workbook's sheet structure.
for c in ws[1]:
    c.font = Font(name="Aptos", bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="234B6B")
    c.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[1].height = 34

row = 2
for month_index in range(36):
    year = 2023 + month_index // 12
    month = month_index % 12 + 1
    # Retail demand tends to rise around the festive period and dip after it.
    season = {1: -3, 2: -1, 3: 0, 4: 1, 5: 0, 6: -1, 7: 0, 8: 1, 9: 2, 10: 7, 11: 9, 12: 5}[month]
    for d in dealers:
        year_offset = year - 2023
        target = max(62, min(128, d["target"] + season * 0.65 + year_offset * rng.uniform(-0.4, 1.4) + rng.gauss(0, 5.5)))
        growth = max(-22, min(35, d["growth"] + year_offset * rng.uniform(-0.4, 1.8) + season * 0.28 + rng.gauss(0, 4.0)))
        inventory = max(18, min(210, d["inventory"] - year_offset * rng.uniform(0, 3) - season * 1.3 + rng.gauss(0, 12)))
        delay = max(0, min(35, d["delay"] - year_offset * 0.25 + rng.gauss(0, 3.0)))
        sla = max(68, min(100, d["sla"] + year_offset * 0.4 + rng.gauss(0, 3.0)))
        complaints = max(0.1, min(12, d["complaints"] + (100 - sla) * 0.045 + rng.gauss(0, 0.65)))
        potential = max(30, round(d["potential"] * (1 + year_offset * 0.018) + rng.gauss(0, d["potential"] * 0.045)))
        frt = max(0.3, min(24, d["frt"] - year_offset * 0.08 + (100 - sla) * 0.035 + rng.gauss(0, 0.7)))
        art = max(2, min(120, d["art"] - year_offset * 0.45 + (100 - sla) * 0.22 + rng.gauss(0, 3.5)))
        satisfaction = max(2.5, min(5, d["satisfaction"] + (sla - d["sla"]) * 0.012 - complaints * 0.018 + rng.gauss(0, 0.16)))
        values = [d["id"], d["name"], f"{year:04d}-{month:02d}", d["state"], d["region"],
                  round(target, 1), round(growth, 1), round(inventory), round(delay, 1),
                  round(sla, 1), round(complaints, 2), potential, round(frt, 1), round(art, 1), round(satisfaction, 2)]
        for col, value in enumerate(values, start=1):
            ws.cell(row, col, value)
        row += 1

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:O{row-1}"
widths = [13, 26, 12, 20, 14, 22, 20, 23, 23, 25, 37, 30, 14, 14, 34]
for col, width in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(col)].width = width

# Keep data clearly identified as a synthetic demonstration in the existing Read Me.
readme = wb["Read Me"]
readme["B13"] = "All 1,620 dealer-month rows (45 fictional dealers, Jan 2023–Dec 2025) are synthetic demonstration data. Do not treat them as real dealer observations."
readme.column_dimensions["B"].width = 110
readme["B13"].alignment = Alignment(wrap_text=True, vertical="top")
readme.row_dimensions[13].height = 32

wb.save(OUTPUT)
print(OUTPUT)
print(f"Rows: {ws.max_row - 1}; dealers: {len(dealers)}; months: 2023-01 through 2025-12")
